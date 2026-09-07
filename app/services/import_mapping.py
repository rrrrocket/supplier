from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from difflib import SequenceMatcher
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl.utils import column_index_from_string, get_column_letter


MAX_IMPORT_ROWS = 10_000

FIELD_DEFINITIONS: dict[str, dict[str, Any]] = {
    "product_name": {
        "label": "商品名称",
        "required": True,
        "aliases": [
            "product_name",
            "product name",
            "item name",
            "商品名称",
            "产品名称",
            "货品名称",
            "品名",
            "名称",
        ],
    },
    "brand": {
        "label": "品牌",
        "required": False,
        "aliases": ["brand", "品牌", "牌子", "品牌名称"],
    },
    "model": {
        "label": "型号",
        "required": False,
        "aliases": ["model", "型号", "规格型号", "款号", "model no"],
    },
    "category": {
        "label": "类目",
        "required": True,
        "aliases": ["category", "类目", "分类", "品类", "产品类别", "商品分类"],
    },
    "supplier_sku": {
        "label": "供应商 SKU",
        "required": True,
        "aliases": [
            "supplier_sku",
            "supplier sku",
            "sku",
            "货号",
            "商品编码",
            "产品编码",
            "物料编码",
            "item no",
            "item number",
            "article no",
        ],
    },
    "price": {
        "label": "采购价",
        "required": True,
        "aliases": [
            "price",
            "cost",
            "unit price",
            "采购价",
            "供货价",
            "成本价",
            "批发价",
            "出厂价",
            "含税价",
            "未税价",
            "不含税价",
            "单价",
        ],
    },
    "currency": {
        "label": "币种",
        "required": False,
        "aliases": ["currency", "currency code", "币种", "货币", "结算币种"],
    },
    "moq": {
        "label": "起订量",
        "required": False,
        "aliases": ["moq", "minimum order", "min order qty", "起订量", "最小起订量"],
    },
    "stock_qty": {
        "label": "库存",
        "required": False,
        "aliases": [
            "stock_qty",
            "stock",
            "inventory",
            "available qty",
            "库存",
            "库存数量",
            "可售库存",
            "可售数",
            "现货数量",
        ],
    },
    "lead_time_days": {
        "label": "交期天数",
        "required": False,
        "aliases": [
            "lead_time_days",
            "lead time",
            "delivery days",
            "交期",
            "交期天数",
            "交货周期",
            "生产周期",
            "发货天数",
        ],
    },
    "fulfillment_mode": {
        "label": "履约模式",
        "required": False,
        "aliases": ["fulfillment_mode", "fulfillment", "履约模式", "合作模式", "供货方式"],
    },
    "status": {
        "label": "状态",
        "required": False,
        "aliases": ["status", "状态", "商品状态", "报价状态", "是否在售"],
    },
}

FIELD_ORDER = list(FIELD_DEFINITIONS)
DEFAULT_VALUES = {
    "currency": "CNY",
    "moq": "1",
    "stock_qty": "0",
    "lead_time_days": "3",
    "fulfillment_mode": "PURCHASE",
    "status": "ACTIVE",
}

PRICE_NEGATIVE_WORDS = ("零售", "市场价", "建议价", "吊牌价", "销售价")
FIELD_NEGATIVE_WORDS = {
    "product_name": ("品牌", "型号", "分类", "类目", "编码", "货号", "SKU"),
    "brand": ("商品", "产品", "型号", "编码"),
    "model": ("商品", "产品", "品牌", "编码"),
    "supplier_sku": ("价格", "库存", "类目", "分类"),
}


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[dict[str, str]]
    header_row: int
    sheet_name: str | None = None


@dataclass(frozen=True)
class WorkbookColumn:
    key: str
    header: str
    label: str


@dataclass
class WorkbookSheet:
    name: str
    index: int
    estimated_rows: int
    column_count: int
    suggested_header_row: int
    columns: list[WorkbookColumn]
    sample_rows: list[dict[str, str]]
    suggested_mapping: dict[str, str]


@dataclass
class WorkbookInspection:
    file_name: str
    active_sheet: str
    sheets: list[WorkbookSheet]


@dataclass(frozen=True)
class SheetImportConfig:
    sheet_name: str
    header_row: int
    mapping: dict[str, str]
    defaults: dict[str, str]


@dataclass
class SourcedRow:
    values: dict[str, str]
    source_sheet: str
    source_row: int


def _normalize_token(value: str) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", value.strip().lower())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def decode_csv(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文件编码无法识别，请使用 UTF-8 或 GB18030 编码")


def _header_match_score(header: str, field: str) -> float:
    normalized_header = _normalize_token(header)
    if not normalized_header:
        return 0
    if field == "price" and any(word in normalized_header for word in PRICE_NEGATIVE_WORDS):
        return 0
    if any(
        _normalize_token(word) in normalized_header
        for word in FIELD_NEGATIVE_WORDS.get(field, ())
    ):
        return 0

    best = 0.0
    for alias in FIELD_DEFINITIONS[field]["aliases"]:
        normalized_alias = _normalize_token(alias)
        if normalized_header == normalized_alias:
            return 100.0
        if len(normalized_alias) >= 2 and normalized_alias in normalized_header:
            best = max(best, 88.0)
        if len(normalized_header) >= 2 and normalized_header in normalized_alias:
            best = max(best, 82.0)
        ratio = SequenceMatcher(None, normalized_header, normalized_alias).ratio()
        if ratio >= 0.72:
            best = max(best, ratio * 80)
    return best


def _row_header_score(row: list[Any]) -> float:
    cells = [_cell_text(value) for value in row]
    nonempty = [value for value in cells if value]
    if len(nonempty) < 2:
        return -1
    recognized = 0
    for value in nonempty:
        if max(_header_match_score(value, field) for field in FIELD_ORDER) >= 82:
            recognized += 1
    numeric = sum(bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value)) for value in nonempty)
    return recognized * 20 + len(nonempty) - numeric * 3


def _description_header_row(description: str) -> int | None:
    match = re.search(r"第\s*(\d+)\s*行(?:是|为|作为)?表头", description)
    if not match:
        return None
    row_number = int(match.group(1))
    return row_number if row_number > 0 else None


def _deduplicate_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        name = _cell_text(value) or f"未命名列{index}"
        counts[name] = counts.get(name, 0) + 1
        headers.append(name if counts[name] == 1 else f"{name} #{counts[name]}")
    return headers


def _build_table(
    matrix: list[list[Any]],
    description: str,
    sheet_name: str | None = None,
) -> ParsedTable:
    matrix = [list(row) for row in matrix]
    if not any(any(_cell_text(value) for value in row) for row in matrix):
        raise ValueError("表格中没有可读取的数据")

    override = _description_header_row(description)
    if override is not None:
        header_index = override - 1
        if header_index >= len(matrix):
            raise ValueError(f"描述指定第 {override} 行为表头，但文件没有这一行")
    else:
        candidates = matrix[: min(20, len(matrix))]
        header_index = max(range(len(candidates)), key=lambda index: _row_header_score(candidates[index]))

    headers = _deduplicate_headers(matrix[header_index])
    rows: list[dict[str, str]] = []
    for raw_row in matrix[header_index + 1 :]:
        values = list(raw_row) + [None] * max(0, len(headers) - len(raw_row))
        row = {header: _cell_text(values[index]) for index, header in enumerate(headers)}
        if any(row.values()):
            rows.append(row)
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValueError(f"单次最多导入 {MAX_IMPORT_ROWS} 行，请拆分文件后重试")
    if not rows:
        raise ValueError("表头下方没有商品数据")
    return ParsedTable(headers=headers, rows=rows, header_row=header_index + 1, sheet_name=sheet_name)


def _read_csv(raw: bytes, description: str) -> ParsedTable:
    text = decode_csv(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    matrix = list(csv.reader(StringIO(text), dialect))
    return _build_table(matrix, description)


def _preferred_sheet_name(sheet_names: list[str], description: str) -> str | None:
    normalized_description = _normalize_token(description)
    for name in sheet_names:
        if _normalize_token(name) in normalized_description:
            return name
    return None


def _load_xlsx_workbook(raw: bytes, *, data_only: bool):
    from openpyxl import load_workbook

    try:
        return load_workbook(BytesIO(raw), read_only=True, data_only=data_only)
    except TypeError as error:
        if "Fill" not in str(error):
            raise

        with ZipFile(BytesIO(raw)) as incoming:
            styles = incoming.read("xl/styles.xml")
            if re.search(rb"<fill\s*/>", styles) is None:
                raise

            output = BytesIO()
            with ZipFile(output, "w") as outgoing:
                for info in incoming.infolist():
                    payload = incoming.read(info.filename)
                    if info.filename == "xl/styles.xml":
                        payload = re.sub(
                            rb"<fill\s*/>",
                            b'<fill><patternFill patternType="none"/></fill>',
                            payload,
                        )
                    outgoing.writestr(info, payload)

        return load_workbook(BytesIO(output.getvalue()), read_only=True, data_only=data_only)


def _suggested_column_mapping(columns: list[WorkbookColumn]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used_keys: set[str] = set()
    priority = ["product_name", "supplier_sku", "price", "category"] + [
        field
        for field in FIELD_ORDER
        if field not in {"product_name", "supplier_sku", "price", "category"}
    ]
    for field in priority:
        ranked = sorted(
            (
                (_header_match_score(column.header, field), column.key)
                for column in columns
                if column.header and column.key not in used_keys
            ),
            reverse=True,
        )
        if ranked and ranked[0][0] >= 72:
            mapping[field] = ranked[0][1]
            used_keys.add(ranked[0][1])
    return mapping


def _inspection_header_score(row: list[Any]) -> float:
    cells = [_cell_text(value) for value in row]
    nonempty = [value for value in cells if value]
    if not nonempty:
        return -1
    recognized = sum(
        max(_header_match_score(value, field) for field in FIELD_ORDER) >= 82
        for value in nonempty
    )
    numeric = sum(bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value)) for value in nonempty)
    return recognized * 20 + len(nonempty) - numeric * 3


def _inspect_workbook_sheet(
    matrix: list[list[Any]],
    *,
    name: str,
    index: int,
) -> WorkbookSheet:
    column_count = max((len(row) for row in matrix), default=0)
    candidates = matrix[: min(20, len(matrix))]
    header_index = (
        max(
            range(len(candidates)),
            key=lambda row_index: _inspection_header_score(candidates[row_index]),
        )
        if candidates
        else 0
    )
    header_values = matrix[header_index] if matrix else []
    headers = [
        _cell_text(header_values[column_index])
        if column_index < len(header_values)
        else ""
        for column_index in range(column_count)
    ]
    columns = []
    for column_index, header in enumerate(headers):
        key = get_column_letter(column_index + 1)
        label = f"{key}列 · {header}" if header else f"{key}列（无表头）"
        columns.append(WorkbookColumn(key=key, header=header, label=label))

    data_rows: list[tuple[int, list[Any]]] = []
    for row_index, row in enumerate(matrix[header_index + 1 :], start=header_index + 2):
        if any(_cell_text(value) for value in row):
            data_rows.append((row_index, row))

    sample_rows = []
    for _, row in data_rows[:5]:
        sample_rows.append(
            {
                get_column_letter(column_index + 1): _cell_text(
                    row[column_index] if column_index < len(row) else None
                )
                for column_index in range(column_count)
            }
        )

    return WorkbookSheet(
        name=name,
        index=index,
        estimated_rows=len(data_rows),
        column_count=column_count,
        suggested_header_row=header_index + 1,
        columns=columns,
        sample_rows=sample_rows,
        suggested_mapping=_suggested_column_mapping(columns),
    )


def inspect_excel_workbook(raw: bytes, filename: str) -> WorkbookInspection:
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        workbook = _load_xlsx_workbook(raw, data_only=True)
        active_sheet = workbook.active.title
        sheets = [
            _inspect_workbook_sheet(
                [list(row) for row in sheet.iter_rows(values_only=True)],
                name=sheet.title,
                index=index,
            )
            for index, sheet in enumerate(workbook.worksheets)
        ]
    elif extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=raw)
        active_index = workbook.sheet_active
        active_sheet = workbook.sheet_by_index(active_index).name
        sheets = [
            _inspect_workbook_sheet(
                [sheet.row_values(row_index) for row_index in range(sheet.nrows)],
                name=sheet.name,
                index=index,
            )
            for index, sheet in enumerate(workbook.sheets())
        ]
    else:
        raise ValueError("仅支持 XLSX 和 XLS 工作簿")

    return WorkbookInspection(
        file_name=filename,
        active_sheet=active_sheet,
        sheets=sheets,
    )


def _validate_sheet_configs(
    configs: list[SheetImportConfig],
    sheet_names: list[str],
    sheet_dimensions: dict[str, tuple[int, int]],
) -> dict[str, dict[str, int]]:
    seen: set[str] = set()
    column_indexes: dict[str, dict[str, int]] = {}
    for config in configs:
        if config.sheet_name in seen:
            raise ValueError(f"工作表不能重复选择：{config.sheet_name}")
        seen.add(config.sheet_name)
        if config.sheet_name not in sheet_names:
            raise ValueError(f"工作表不存在：{config.sheet_name}")

        row_count, column_count = sheet_dimensions[config.sheet_name]
        if config.header_row < 1 or config.header_row > row_count:
            raise ValueError(
                f"工作表 {config.sheet_name} 不存在第 {config.header_row} 行表头"
            )

        indexes: dict[str, int] = {}
        for field, source in config.mapping.items():
            if field not in FIELD_ORDER:
                continue
            if re.fullmatch(r"[A-Z]+", source) is None:
                raise ValueError(
                    f"工作表 {config.sheet_name} 的字段映射引用了无效列：{source}"
                )
            column_index = column_index_from_string(source)
            if column_index > column_count:
                raise ValueError(
                    f"工作表 {config.sheet_name} 的字段映射引用了无效列：{source}"
                )
            indexes[field] = column_index - 1
        column_indexes[config.sheet_name] = indexes
    return column_indexes


def _sourced_values(
    row: list[Any],
    column_indexes: dict[str, int],
    defaults: dict[str, str],
) -> dict[str, str]:
    values: dict[str, str] = {}
    for field in FIELD_ORDER:
        column_index = column_indexes.get(field)
        value = (
            _cell_text(row[column_index])
            if column_index is not None and column_index < len(row)
            else ""
        )
        values[field] = value or defaults.get(field, "")
    return values


def parse_excel_sheets(
    raw: bytes,
    filename: str,
    configs: list[SheetImportConfig],
) -> list[SourcedRow]:
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        workbook = _load_xlsx_workbook(raw, data_only=True)
        sheet_names = workbook.sheetnames
        sheet_dimensions = {
            name: (workbook[name].max_row, workbook[name].max_column)
            for name in {config.sheet_name for config in configs}
            if name in sheet_names
        }
    elif extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=raw)
        sheet_names = workbook.sheet_names()
        sheet_dimensions = {
            name: (
                workbook.sheet_by_name(name).nrows,
                workbook.sheet_by_name(name).ncols,
            )
            for name in {config.sheet_name for config in configs}
            if name in sheet_names
        }
    else:
        raise ValueError("仅支持 XLSX 和 XLS 工作簿")

    column_indexes = _validate_sheet_configs(configs, sheet_names, sheet_dimensions)
    sourced_rows: list[SourcedRow] = []
    for config in configs:
        defaults = {**DEFAULT_VALUES, **config.defaults}
        if extension == ".xlsx":
            sheet = workbook[config.sheet_name]
            rows = enumerate(
                sheet.iter_rows(min_row=config.header_row + 1, values_only=True),
                start=config.header_row + 1,
            )
        else:
            sheet = workbook.sheet_by_name(config.sheet_name)
            rows = (
                (row_index + 1, sheet.row_values(row_index))
                for row_index in range(config.header_row, sheet.nrows)
            )

        for source_row, raw_row in rows:
            row = list(raw_row)
            if not any(_cell_text(value) for value in row):
                continue
            sourced_rows.append(
                SourcedRow(
                    values=_sourced_values(
                        row,
                        column_indexes[config.sheet_name],
                        defaults,
                    ),
                    source_sheet=config.sheet_name,
                    source_row=source_row,
                )
            )
            if len(sourced_rows) > MAX_IMPORT_ROWS:
                raise ValueError(
                    f"单次最多导入 {MAX_IMPORT_ROWS} 行，请拆分文件后重试"
                )
    return sourced_rows


def _read_xlsx(raw: bytes, description: str) -> ParsedTable:
    workbook = _load_xlsx_workbook(raw, data_only=True)
    preferred = _preferred_sheet_name(workbook.sheetnames, description)
    sheets = [workbook[preferred]] if preferred else list(workbook.worksheets)
    candidates: list[tuple[float, ParsedTable]] = []
    for sheet in sheets:
        matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
        try:
            table = _build_table(matrix, description, sheet.title)
        except ValueError:
            continue
        score = max(_row_header_score(row) for row in matrix[: min(20, len(matrix))])
        candidates.append((score, table))
    if not candidates:
        raise ValueError("Excel 中没有可读取的商品数据")
    return max(candidates, key=lambda item: (item[0], len(item[1].rows)))[1]


def _read_xls(raw: bytes, description: str) -> ParsedTable:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=raw)
    sheet_names = workbook.sheet_names()
    preferred = _preferred_sheet_name(sheet_names, description)
    sheets = [workbook.sheet_by_name(preferred)] if preferred else workbook.sheets()
    candidates: list[tuple[float, ParsedTable]] = []
    for sheet in sheets:
        matrix = [sheet.row_values(index) for index in range(sheet.nrows)]
        try:
            table = _build_table(matrix, description, sheet.name)
        except ValueError:
            continue
        score = max(_row_header_score(row) for row in matrix[: min(20, len(matrix))])
        candidates.append((score, table))
    if not candidates:
        raise ValueError("Excel 中没有可读取的商品数据")
    return max(candidates, key=lambda item: (item[0], len(item[1].rows)))[1]


def _pdf_table_score(table: ParsedTable) -> float:
    recognized = sum(
        max(_header_match_score(header, field) for field in FIELD_ORDER) >= 72
        for header in table.headers
    )
    return recognized * 100 + len(table.rows)


def _header_signature(headers: list[str]) -> tuple[str, ...]:
    return tuple(_normalize_token(header.split(" #", 1)[0]) for header in headers)


def _read_pdf(raw: bytes, description: str) -> ParsedTable:
    import pdfplumber

    candidates: list[tuple[float, int, ParsedTable]] = []
    has_text = False
    with pdfplumber.open(BytesIO(raw)) as document:
        if len(document.pages) > 100:
            raise ValueError("单个 PDF 最多支持 100 页，请拆分后重试")
        for page_number, page in enumerate(document.pages, start=1):
            page_tables = page.extract_tables()
            extracted_count = len(candidates)
            for matrix in page_tables:
                if not matrix:
                    continue
                try:
                    table = _build_table(matrix, description, f"PDF 第 {page_number} 页")
                except ValueError:
                    continue
                candidates.append((_pdf_table_score(table), page_number, table))

            text = page.extract_text(layout=True) or ""
            if text.strip():
                has_text = True
            if len(candidates) > extracted_count or not text.strip():
                continue
            matrix = []
            for line in text.splitlines():
                cells = [cell.strip() for cell in re.split(r"\s{2,}|\t+", line.strip())]
                if len([cell for cell in cells if cell]) >= 2:
                    matrix.append(cells)
            if not matrix:
                continue
            try:
                table = _build_table(matrix, description, f"PDF 第 {page_number} 页")
            except ValueError:
                continue
            candidates.append((_pdf_table_score(table), page_number, table))

    if not candidates:
        if not has_text:
            raise ValueError("PDF 是纯扫描图片，当前无法可靠读取价格，请先进行 OCR 后重试")
        raise ValueError("PDF 中未识别到结构化表格，请补充说明表头和字段对应关系")

    _, _, best = max(candidates, key=lambda item: item[0])
    signature = _header_signature(best.headers)
    matching = [item for item in candidates if _header_signature(item[2].headers) == signature]
    matching.sort(key=lambda item: item[1])
    merged_rows: list[dict[str, str]] = []
    page_numbers: list[int] = []
    for _, page_number, table in matching:
        merged_rows.extend(table.rows)
        page_numbers.append(page_number)
        if len(merged_rows) > MAX_IMPORT_ROWS:
            raise ValueError(f"单次最多导入 {MAX_IMPORT_ROWS} 行，请拆分文件后重试")

    page_label = "、".join(str(page) for page in sorted(set(page_numbers)))
    return ParsedTable(
        headers=best.headers,
        rows=merged_rows,
        header_row=best.header_row,
        sheet_name=f"PDF 第 {page_label} 页",
    )


def parse_table(raw: bytes, filename: str, description: str = "") -> ParsedTable:
    extension = Path(filename).suffix.lower()
    try:
        if extension in {".csv", ".tsv", ".txt"}:
            return _read_csv(raw, description)
        if extension == ".xlsx":
            return _read_xlsx(raw, description)
        if extension == ".xls":
            return _read_xls(raw, description)
        if extension == ".pdf":
            return _read_pdf(raw, description)
        raise ValueError("支持 CSV、XLSX、XLS 和 PDF 格式")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("文件无法解析，请确认表格未损坏且未加密") from exc


def _description_mapping(headers: list[str], description: str) -> dict[str, str]:
    result: dict[str, str] = {}
    clauses = re.split(r"[，,；;。\n]+", description)
    for clause in clauses:
        normalized_clause = _normalize_token(clause)
        if not normalized_clause:
            continue
        source_headers = [
            header for header in headers if _normalize_token(header) in normalized_clause
        ]
        if not source_headers:
            continue
        for field, definition in FIELD_DEFINITIONS.items():
            target_found = any(
                _normalize_token(alias) in normalized_clause for alias in definition["aliases"]
            )
            if not target_found:
                continue
            source = max(source_headers, key=lambda header: len(_normalize_token(header)))
            result[field] = source
    return result


def infer_mapping(headers: list[str], description: str = "") -> dict[str, str]:
    mapping = _description_mapping(headers, description)
    used_headers = set(mapping.values())
    priority = ["product_name", "supplier_sku", "price", "category"] + [
        field for field in FIELD_ORDER if field not in {"product_name", "supplier_sku", "price", "category"}
    ]
    for field in priority:
        if field in mapping:
            continue
        ranked = sorted(
            (
                (_header_match_score(header, field), header)
                for header in headers
                if header not in used_headers
            ),
            reverse=True,
        )
        if ranked and ranked[0][0] >= 72:
            mapping[field] = ranked[0][1]
            used_headers.add(ranked[0][1])
    return mapping


def infer_defaults(description: str = "") -> dict[str, str]:
    defaults = dict(DEFAULT_VALUES)
    normalized = description.upper()
    currency_words = {
        "人民币": "CNY",
        "RMB": "CNY",
        "CNY": "CNY",
        "美元": "USD",
        "USD": "USD",
        "欧元": "EUR",
        "EUR": "EUR",
        "日元": "JPY",
        "JPY": "JPY",
        "港币": "HKD",
        "HKD": "HKD",
        "英镑": "GBP",
        "GBP": "GBP",
    }
    for word, code in currency_words.items():
        if word in normalized:
            defaults["currency"] = code
            break

    category_match = re.search(
        r"(?:全部|所有|统一)(?:商品)?(?:的)?(?:类目|分类|品类)"
        r"(?:统一)?(?:填(?:写|成)?|设为|为|是)?[:：]?\s*([^，,。；;\n]+)",
        description,
    )
    if category_match:
        defaults["category"] = category_match.group(1).strip()

    numeric_patterns = {
        "moq": r"(?:起订量|MOQ)(?:默认|统一)?(?:填|为|是|按)?[:：]?\s*(\d+)",
        "stock_qty": r"(?:库存)(?:默认|空值)?(?:填|为|是|按)?[:：]?\s*(\d+)",
        "lead_time_days": r"(?:交期)(?:默认|统一)?(?:填|为|是|按)?[:：]?\s*(\d+)",
    }
    for field, pattern in numeric_patterns.items():
        match = re.search(pattern, description, re.IGNORECASE)
        if match:
            defaults[field] = match.group(1)

    if "一件代发" in description:
        defaults["fulfillment_mode"] = "DROPSHIP"
    elif "寄售" in description:
        defaults["fulfillment_mode"] = "CONSIGNMENT"
    elif "联营" in description:
        defaults["fulfillment_mode"] = "JOINT_OPERATION"
    return defaults


def normalize_rows(
    table: ParsedTable,
    mapping: dict[str, str],
    defaults: dict[str, str] | None = None,
) -> list[dict[str, str]]:
    effective_defaults = {**DEFAULT_VALUES, **(defaults or {})}
    normalized: list[dict[str, str]] = []
    for raw_row in table.rows:
        row: dict[str, str] = {}
        for field in FIELD_ORDER:
            source = mapping.get(field)
            value = raw_row.get(source, "").strip() if source else ""
            row[field] = value or effective_defaults.get(field, "")
        normalized.append(row)
    return normalized


def mapping_view(mapping: dict[str, str], headers: list[str]) -> list[dict[str, Any]]:
    result = []
    for field, definition in FIELD_DEFINITIONS.items():
        source = mapping.get(field, "")
        score = _header_match_score(source, field) if source else 0
        result.append(
            {
                "field": field,
                "label": definition["label"],
                "required": definition["required"],
                "source": source,
                "confidence": "high" if score >= 95 else "medium" if score >= 80 else "low",
                "available_sources": headers,
            }
        )
    return result
