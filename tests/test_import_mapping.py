from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

import pytest
import xlwt
from openpyxl import Workbook

from app.services.import_mapping import (
    MAX_IMPORT_ROWS,
    SheetImportConfig,
    inspect_excel_workbook,
    parse_excel_sheets,
)


def workbook_bytes_with_empty_fill() -> bytes:
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Sheet1"
    sheet1.append(
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "成本",
        ]
    )
    sheet1.append(
        [
            None,
            None,
            "001",
            None,
            "商品一",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "TRU",
            "1:35",
            "模型",
            None,
            70,
        ]
    )
    sheet2 = workbook.create_sheet("Sheet2")
    sheet2.append(["名称", "编号", "销售价"])
    sheet2.append(["参考商品", "REF-1", 100])
    source = BytesIO()
    workbook.save(source)

    incoming = ZipFile(BytesIO(source.getvalue()))
    output = BytesIO()
    outgoing = ZipFile(output, "w")
    for info in incoming.infolist():
        payload = incoming.read(info.filename)
        if info.filename == "xl/styles.xml":
            payload = payload.replace(
                b'<fills count="2">',
                b'<fills count="3">',
                1,
            )
            payload = payload.replace(b"</fills>", b"<fill/></fills>", 1)
        outgoing.writestr(info, payload)
    outgoing.close()
    incoming.close()
    return output.getvalue()


def workbook_bytes_with_row_counts(*row_counts: int) -> bytes:
    workbook = Workbook()
    for sheet_index, row_count in enumerate(row_counts, start=1):
        sheet = workbook.active if sheet_index == 1 else workbook.create_sheet()
        sheet.title = f"Data{sheet_index}"
        sheet.append(["名称"])
        for row_index in range(row_count):
            sheet.append([f"商品-{sheet_index}-{row_index}"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def workbook_bytes_xls(*, active_second_sheet: bool) -> bytes:
    workbook = xlwt.Workbook()
    first = workbook.add_sheet("First")
    first.write(0, 0, "名称")
    first.write(1, 0, "首表商品")
    second = workbook.add_sheet("报价")
    second.write(0, 0, "报价说明")
    second.write(1, 0, "商品名称")
    second.write(1, 2, "采购价")
    second.write(2, 0, "旧格式商品")
    second.write(2, 2, 88)
    if active_second_sheet:
        second.sheet_visible = True
        second.selected = True
        workbook.set_active_sheet(1)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_inspect_excel_workbook_repairs_empty_fill_and_preserves_active_sheet():
    raw = workbook_bytes_with_empty_fill()
    original_hash = sha256(raw).hexdigest()

    inspection = inspect_excel_workbook(raw, "supplier.xlsx")

    assert sha256(raw).hexdigest() == original_hash
    assert inspection.active_sheet == "Sheet1"
    assert [sheet.name for sheet in inspection.sheets] == ["Sheet1", "Sheet2"]
    assert inspection.sheets[0].columns[2].key == "C"
    assert inspection.sheets[0].columns[2].label == "C列（无表头）"
    assert inspection.sheets[0].columns[17].label == "R列 · 成本"


def test_inspect_xls_uses_visible_sheet_as_active_sheet():
    inspection = inspect_excel_workbook(
        workbook_bytes_xls(active_second_sheet=True),
        "supplier.xls",
    )

    assert inspection.active_sheet == "报价"
    assert [sheet.name for sheet in inspection.sheets] == ["First", "报价"]


def test_inspect_xls_falls_back_to_first_sheet_when_none_is_visible():
    inspection = inspect_excel_workbook(
        workbook_bytes_xls(active_second_sheet=False),
        "supplier.xls",
    )

    assert inspection.active_sheet == "First"


def test_parse_excel_sheets_uses_column_keys_and_preserves_source_coordinates():
    rows = parse_excel_sheets(
        workbook_bytes_with_empty_fill(),
        "supplier.xlsx",
        [
            SheetImportConfig(
                sheet_name="Sheet1",
                header_row=1,
                mapping={
                    "supplier_sku": "C",
                    "product_name": "E",
                    "brand": "N",
                    "model": "O",
                    "category": "P",
                    "price": "R",
                },
                defaults={"currency": "CNY"},
            )
        ],
    )

    assert len(rows) == 1
    assert rows[0].source_sheet == "Sheet1"
    assert rows[0].source_row == 2
    assert rows[0].values["supplier_sku"] == "001"
    assert rows[0].values["product_name"] == "商品一"
    assert rows[0].values["price"] == "70"


def test_parse_xls_sheet_uses_mapping_defaults_and_source_coordinates():
    rows = parse_excel_sheets(
        workbook_bytes_xls(active_second_sheet=True),
        "supplier.xls",
        [
            SheetImportConfig(
                sheet_name="报价",
                header_row=2,
                mapping={"product_name": "A", "price": "C"},
                defaults={"currency": "USD", "category": "模型"},
            )
        ],
    )

    assert len(rows) == 1
    assert rows[0].source_sheet == "报价"
    assert rows[0].source_row == 3
    assert rows[0].values["product_name"] == "旧格式商品"
    assert rows[0].values["price"] == "88"
    assert rows[0].values["currency"] == "USD"
    assert rows[0].values["category"] == "模型"


def test_parse_excel_sheets_rejects_duplicate_sheet_names():
    config = SheetImportConfig("Sheet1", 1, {"price": "R"}, {})

    with pytest.raises(ValueError, match="^工作表不能重复选择：Sheet1$"):
        parse_excel_sheets(
            workbook_bytes_with_empty_fill(),
            "supplier.xlsx",
            [config, config],
        )


def test_parse_excel_sheets_rejects_missing_sheet_names():
    with pytest.raises(ValueError, match="^工作表不存在：Missing$"):
        parse_excel_sheets(
            workbook_bytes_with_empty_fill(),
            "supplier.xlsx",
            [SheetImportConfig("Missing", 1, {"product_name": "A"}, {})],
        )


def test_parse_excel_sheets_rejects_missing_header_rows():
    with pytest.raises(ValueError, match="^工作表 Sheet2 不存在第 3 行表头$"):
        parse_excel_sheets(
            workbook_bytes_with_empty_fill(),
            "supplier.xlsx",
            [SheetImportConfig("Sheet2", 3, {"product_name": "A"}, {})],
        )


def test_parse_excel_sheets_rejects_invalid_column_letters():
    with pytest.raises(ValueError, match="^工作表 Sheet2 的字段映射引用了无效列：A1$"):
        parse_excel_sheets(
            workbook_bytes_with_empty_fill(),
            "supplier.xlsx",
            [SheetImportConfig("Sheet2", 1, {"product_name": "A1"}, {})],
        )


def test_parse_excel_sheets_enforces_combined_row_limit():
    raw = workbook_bytes_with_row_counts(MAX_IMPORT_ROWS // 2, MAX_IMPORT_ROWS // 2 + 1)

    with pytest.raises(
        ValueError,
        match="^单次最多导入 10000 行，请拆分文件后重试$",
    ):
        parse_excel_sheets(
            raw,
            "supplier.xlsx",
            [
                SheetImportConfig("Data1", 1, {"product_name": "A"}, {}),
                SheetImportConfig("Data2", 1, {"product_name": "A"}, {}),
            ],
        )
