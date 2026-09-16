from __future__ import annotations

import asyncio
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
import xlwt
from fastapi.testclient import TestClient
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from sqlalchemy import select

import app.services.import_mapping as import_mapping
import app.api.routes.imports as import_routes
from app.api.routes.imports import MAX_IMPORT_FORM_PART_SIZE
from app.db.session import SessionLocal
from app.models.entities import (
    Brand,
    CatalogStatus,
    CommercialMode,
    ImportJob,
    ImportStatus,
    InventorySnapshot,
    Organization,
    OrganizationType,
    Product,
    SupplierBrandCooperation,
    SupplierOffer,
    SupplierSku,
    User,
)
from app.services.list_pagination import decode_list_cursor
from tests.conftest import SUPPLIER_EMAIL, SUPPLIER_PASSWORD


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_import_integer_grammar_matches_shared_browser_vectors() -> None:
    vectors = json.loads(
        (Path(__file__).with_name("import_integer_vectors.json")).read_text()
    )
    for vector in vectors:
        if vector.get("error"):
            with pytest.raises(ValueError) as exc_info:
                import_routes._to_int(vector["raw"], 0, "库存")
            expected = (
                "库存不能小于0"
                if vector["error"] == "negative"
                else (
                    "库存不能大于2147483647"
                    if vector["error"] == "range"
                    else "库存必须是整数"
                )
            )
            assert str(exc_info.value) == expected, vector["raw"]
        else:
            assert import_routes._to_int(vector["raw"], 0, "库存") == vector["parsed"]


def two_sheet_workbook_content() -> bytes:
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Sheet1"
    sheet1.append([None, None, "货号"])
    sheet1.append([None, None, "SKU-001"])
    sheet2 = workbook.create_sheet("Sheet2")
    sheet2.append(["名称"])
    sheet2.append(["商品二"])
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def multi_sheet_offer_workbook_content(second_sku: str = "DUP-001") -> bytes:
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Sheet1"
    sheet1.append(["货号", "名称", "成本"])
    sheet1.append(["DUP-001", "商品甲", 10])
    sheet2 = workbook.create_sheet("Sheet2")
    sheet2.append(["供应商成本表"])
    sheet2.append(["产品编码", "产品名称", "出厂价"])
    sheet2.append([second_sku, "商品乙", 12])
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def two_sheet_xls_content() -> bytes:
    workbook = xlwt.Workbook()
    first = workbook.add_sheet("Sheet1")
    first.write(0, 0, "货号")
    first.write(1, 0, "XLS-001")
    second = workbook.add_sheet("Sheet2")
    second.write(0, 0, "名称")
    second.write(1, 0, "商品二")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def xlsx_with_encrypted_entry_flag() -> bytes:
    payload = bytearray(two_sheet_workbook_content())
    for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        position = 0
        while True:
            position = payload.find(signature, position)
            if position < 0:
                break
            offset = position + flag_offset
            flags = int.from_bytes(payload[offset : offset + 2], "little") | 0x1
            payload[offset : offset + 2] = flags.to_bytes(2, "little")
            position += len(signature)
    return bytes(payload)


def truncated_ole_xls_content() -> bytes:
    """Compound-document signature followed by an intentionally truncated header."""
    return bytes.fromhex("d0cf11e0a1b11ae1") + (b"\x00" * 56)


def highly_expanded_zip_content() -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"x" * 700)
        archive.writestr("xl/worksheets/sheet2.xml", b"x" * 700)
    return output.getvalue()


def multi_sheet_configs() -> list[dict[str, object]]:
    defaults = {"brand": "TEST", "category": "测试类目", "currency": "CNY"}
    mapping = {"supplier_sku": "A", "product_name": "B", "price": "C"}
    return [
        {
            "sheet_name": "Sheet1",
            "header_row": 1,
            "mapping": mapping,
            "defaults": defaults,
        },
        {
            "sheet_name": "Sheet2",
            "header_row": 2,
            "mapping": mapping,
            "defaults": defaults,
        },
    ]


def submitted_offer_row(
    *,
    sku: str,
    name: str,
    source_sheet: object,
    source_row: object,
    included: object = True,
    price: str = "10",
) -> dict[str, object]:
    return {
        "source_sheet": source_sheet,
        "source_row": source_row,
        "included": included,
        "values": {
            "product_name": name,
            "brand": "TEST",
            "model": sku,
            "category": "测试类目",
            "supplier_sku": sku,
            "price": price,
            "currency": "CNY",
            "moq": "1",
            "stock_qty": "5",
            "lead_time_days": "3",
            "fulfillment_mode": "PURCHASE",
            "status": "ACTIVE",
        },
    }


def test_health(client: TestClient) -> None:
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_product_page_cursor_does_not_shift_when_an_earlier_row_is_updated(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    with SessionLocal() as db:
        supplier = db.scalar(
            select(Organization).where(Organization.code == "TEST-SUPPLIER")
        )
        brand = db.scalar(select(Brand).where(Brand.code == "TEST-BRAND"))
        assert supplier is not None and brand is not None
        products = [
            Product(
                id=f"page-{suffix}-{index}",
                created_by_organization_id=supplier.id,
                brand_id=brand.id,
                name=f"分页商品-{suffix}-{index}",
                category="测试",
            )
            for index in range(3)
        ]
        db.add_all(products)
        db.commit()

    first = authenticated_client.get(
        "/api/products/page",
        params={"q": f"分页商品-{suffix}", "limit": 2},
    )
    assert first.status_code == 200
    assert len(first.json()["items"]) == 2
    cursor = first.json()["next_cursor"]
    assert cursor is not None
    filters = {
        "organization_id": supplier.id,
        "q": f"分页商品-{suffix}",
        "status": None,
    }
    assert decode_list_cursor(cursor, "products", filters) == (
        first.json()["items"][-1]["id"]
    )
    with pytest.raises(Exception) as exc_info:
        decode_list_cursor(
            cursor,
            "products",
            {**filters, "organization_id": "different-supplier"},
        )
    assert getattr(exc_info.value, "status_code", None) == 400
    with SessionLocal() as db:
        changed = db.get(Product, first.json()["items"][0]["id"])
        assert changed is not None
        changed.updated_at = datetime.now(timezone.utc)
        db.commit()

    second = authenticated_client.get(
        "/api/products/page",
        params={
            "q": f"分页商品-{suffix}",
            "limit": 2,
            "cursor": cursor,
        },
    )
    assert second.status_code == 200
    ids = [item["id"] for item in first.json()["items"] + second.json()["items"]]
    assert ids == [f"page-{suffix}-{index}" for index in range(3)]

    rebound = authenticated_client.get(
        "/api/products/page",
        params={
            "q": f"不同筛选-{suffix}",
            "limit": 2,
            "cursor": first.json()["next_cursor"],
        },
    )
    assert rebound.status_code == 400
    assert rebound.json()["detail"] == "分页游标无效"


def test_offer_page_cursor_does_not_shift_when_an_earlier_row_is_updated(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    with SessionLocal() as db:
        supplier = db.scalar(
            select(Organization).where(Organization.code == "TEST-SUPPLIER")
        )
        brand = db.scalar(select(Brand).where(Brand.code == "TEST-BRAND"))
        assert supplier is not None and brand is not None
        products = [
            Product(
                id=f"offer-product-{suffix}-{index}",
                created_by_organization_id=supplier.id,
                brand_id=brand.id,
                name=f"报价分页商品-{suffix}-{index}",
                category="测试",
            )
            for index in range(3)
        ]
        skus = [
            SupplierSku(
                id=f"offer-sku-{suffix}-{index}",
                supplier_id=supplier.id,
                brand_id=brand.id,
                product_id=products[index].id,
                supplier_sku_code=f"报价分页货号-{suffix}-{index}",
                status=CatalogStatus.ACTIVE.value,
            )
            for index in range(3)
        ]
        offers = [
            SupplierOffer(
                id=f"offer-page-{suffix}-{index}",
                organization_id=supplier.id,
                product_id=products[index].id,
                supplier_sku_id=skus[index].id,
                price="10.0000",
                currency="CNY",
                status="ACTIVE",
            )
            for index in range(3)
        ]
        db.add_all([*products, *skus, *offers])
        db.commit()

    first = authenticated_client.get(
        "/api/offers/page",
        params={"q": suffix, "limit": 2},
    )
    assert first.status_code == 200
    assert len(first.json()["items"]) == 2
    offer_cursor = first.json()["next_cursor"]
    assert offer_cursor is not None
    offer_filters = {
        "organization_id": supplier.id,
        "q": suffix,
        "brand": None,
        "status": None,
    }
    assert decode_list_cursor(offer_cursor, "offers", offer_filters) == (
        first.json()["items"][-1]["id"]
    )
    with pytest.raises(Exception) as exc_info:
        decode_list_cursor(
            offer_cursor,
            "offers",
            {**offer_filters, "organization_id": "different-supplier"},
        )
    assert getattr(exc_info.value, "status_code", None) == 400
    with SessionLocal() as db:
        changed = db.get(SupplierOffer, first.json()["items"][0]["id"])
        assert changed is not None
        changed.updated_at = datetime.now(timezone.utc)
        db.commit()

    second = authenticated_client.get(
        "/api/offers/page",
        params={
            "q": suffix,
            "limit": 2,
            "cursor": offer_cursor,
        },
    )
    assert second.status_code == 200
    ids = [item["id"] for item in first.json()["items"] + second.json()["items"]]
    assert ids == [f"offer-page-{suffix}-{index}" for index in range(3)]


@pytest.mark.parametrize(
    "path",
    ["/api/products", "/api/products/page", "/api/offers", "/api/offers/page"],
)
def test_catalog_search_accepts_full_product_name_length(
    authenticated_client: TestClient,
    path: str,
) -> None:
    response = authenticated_client.get(path, params={"q": "界" * 240})
    assert response.status_code == 200


def test_excel_inspection_returns_all_sheets_and_active_sheet(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/workbook/inspect",
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(two_sheet_workbook_content()),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["active_sheet"] == "Sheet1"
    assert [sheet["name"] for sheet in data["sheets"]] == ["Sheet1", "Sheet2"]
    assert data["sheets"][0]["columns"][2]["key"] == "C"


def test_xls_inspection_returns_real_biff_workbook_metadata(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/workbook/inspect",
        files={
            "file": (
                "supplier.xls",
                BytesIO(two_sheet_xls_content()),
                "application/vnd.ms-excel",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["active_sheet"] == "Sheet1"
    assert [sheet["name"] for sheet in response.json()["sheets"]] == ["Sheet1", "Sheet2"]


def test_excel_inspection_rejects_csv(authenticated_client: TestClient) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/workbook/inspect",
        files={"file": ("supplier.csv", BytesIO(b"sku,name\n1,item\n"), "text/csv")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "工作簿扫描仅支持 XLSX 和 XLS 格式"


def test_multi_sheet_preview_reports_duplicate_sku_conflict(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/preview",
        data={"sheet_configs_json": json.dumps(multi_sheet_configs(), ensure_ascii=False)},
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(multi_sheet_offer_workbook_content()),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total_rows"] == 2
    assert data["conflict_count"] == 1
    assert data["can_import"] is False
    assert {
        (row["source_sheet"], row["source_row"])
        for row in data["preview_rows"]
    } == {("Sheet1", 2), ("Sheet2", 3)}
    assert {row["conflict_group"] for row in data["preview_rows"]} == {"DUP-001"}
    assert all(row["included"] is True for row in data["preview_rows"])


def test_multi_sheet_preview_allows_distinct_skus(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/preview",
        data={"sheet_configs_json": json.dumps(multi_sheet_configs(), ensure_ascii=False)},
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(multi_sheet_offer_workbook_content("UNIQUE-002")),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["can_import"] is True
    assert response.json()["conflict_count"] == 0


def test_multi_sheet_preview_marks_empty_brand_as_needing_correction(
    authenticated_client: TestClient,
) -> None:
    configs = multi_sheet_configs()
    for config in configs:
        config["defaults"] = {"category": "测试类目", "currency": "CNY"}
    response = authenticated_client.post(
        "/api/imports/product-offers/preview",
        data={"sheet_configs_json": json.dumps(configs, ensure_ascii=False)},
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(multi_sheet_offer_workbook_content("UNIQUE-002")),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["can_import"] is False
    assert all("品牌不能为空" in row["errors"] for row in response.json()["preview_rows"])


@pytest.mark.parametrize(
    ("filename", "payload"),
    [
        ("broken.xlsx", b"not a zip workbook"),
        ("broken.xls", b"not a biff workbook"),
        ("broken-ole.xls", truncated_ole_xls_content()),
        ("mismatch.xlsx", two_sheet_xls_content()),
        ("mismatch.xls", two_sheet_workbook_content()),
        ("encrypted.xlsx", xlsx_with_encrypted_entry_flag()),
    ],
)
def test_workbook_inspection_rejects_bad_or_mismatched_content_safely(
    authenticated_client: TestClient,
    filename: str,
    payload: bytes,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/workbook/inspect",
        files={"file": (filename, BytesIO(payload), "application/octet-stream")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == (
        "文件无法解析，请确认工作簿未损坏且未加密"
    )


@pytest.mark.parametrize(
    ("filename", "payload"),
    [
        ("broken.xlsx", b"not a zip workbook"),
        ("broken.xls", b"not a biff workbook"),
        ("broken-ole.xls", truncated_ole_xls_content()),
        ("mismatch.xlsx", two_sheet_xls_content()),
        ("mismatch.xls", two_sheet_workbook_content()),
        ("encrypted.xlsx", xlsx_with_encrypted_entry_flag()),
    ],
)
def test_configured_preview_rejects_bad_or_mismatched_content_safely(
    authenticated_client: TestClient,
    filename: str,
    payload: bytes,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/preview",
        data={"sheet_configs_json": json.dumps(multi_sheet_configs())},
        files={"file": (filename, BytesIO(payload), "application/octet-stream")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == (
        "文件无法解析，请确认工作簿未损坏且未加密"
    )


@pytest.mark.parametrize(
    ("endpoint", "form_data"),
    [
        ("/api/imports/product-offers/workbook/inspect", {}),
        (
            "/api/imports/product-offers/preview",
            {"sheet_configs_json": json.dumps([multi_sheet_configs()[0]])},
        ),
    ],
)
def test_excel_api_rejects_high_zip_expansion(
    authenticated_client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
    endpoint: str,
    form_data: dict[str, str],
) -> None:
    monkeypatch.setattr(import_mapping, "MAX_WORKBOOK_UNCOMPRESSED_BYTES", 1_024)
    monkeypatch.setattr(import_mapping, "MAX_WORKBOOK_SINGLE_ENTRY_BYTES", 1_024)
    monkeypatch.setattr(import_mapping, "MAX_WORKBOOK_COMPRESSION_RATIO", 10_000)

    response = authenticated_client.post(
        endpoint,
        data=form_data,
        files={"file": ("expanded.xlsx", BytesIO(highly_expanded_zip_content()), XLSX_MIME)},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "工作簿解压后大小超过安全限制"


def test_excel_api_rejects_overwide_sheet(authenticated_client: TestClient) -> None:
    workbook = Workbook()
    workbook.active.cell(row=1, column=import_mapping.MAX_WORKBOOK_COLUMNS + 1, value="超宽")
    output = BytesIO()
    workbook.save(output)

    response = authenticated_client.post(
        "/api/imports/product-offers/workbook/inspect",
        files={"file": ("wide.xlsx", BytesIO(output.getvalue()), XLSX_MIME)},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "工作表行列规模超过安全限制：Sheet"


@pytest.mark.parametrize(
    ("sheet_configs_json", "detail"),
    [
        ("{", "工作表配置格式不正确"),
        ("{}", "工作表配置必须是数组"),
        ("[]", "请至少选择一个工作表"),
        (
            json.dumps([multi_sheet_configs()[0], multi_sheet_configs()[0]]),
            "工作表不能重复选择：Sheet1",
        ),
        (
            json.dumps(
                [
                    {
                        **multi_sheet_configs()[0],
                        "sheet_name": "Missing",
                    }
                ]
            ),
            "工作表不存在：Missing",
        ),
        (
            json.dumps([{**multi_sheet_configs()[0], "header_row": 4}]),
            "工作表 Sheet1 不存在第 4 行表头",
        ),
        (
            json.dumps(
                [
                    {
                        **multi_sheet_configs()[0],
                        "mapping": {"product_name": "A1"},
                    }
                ]
            ),
            "工作表 Sheet1 的字段映射引用了无效列：A1",
        ),
    ],
)
def test_multi_sheet_preview_rejects_invalid_configuration(
    authenticated_client: TestClient,
    sheet_configs_json: str,
    detail: str,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers/preview",
        data={"sheet_configs_json": sheet_configs_json},
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(multi_sheet_offer_workbook_content()),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == detail


def test_final_import_rejects_duplicate_included_skus_without_writing_offer(
    authenticated_client: TestClient,
) -> None:
    sku = f"DUP-FINAL-{uuid4().hex[:8]}"
    rows = [
        submitted_offer_row(
            sku=sku,
            name="商品甲",
            source_sheet="Sheet1",
            source_row=2,
        ),
        submitted_offer_row(
            sku=sku,
            name="商品乙",
            source_sheet="Sheet2",
            source_row=3,
        ),
    ]

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={
            "rows_json": json.dumps(rows, ensure_ascii=False),
            "sheet_configs_json": json.dumps(multi_sheet_configs(), ensure_ascii=False),
        },
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(multi_sheet_offer_workbook_content()),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == f"供应商 SKU 在本次导入中重复：{sku}"
    offers = authenticated_client.get("/api/offers", params={"q": sku}).json()
    assert offers == []


def test_final_import_excludes_duplicate_row_and_preserves_partial_error_source(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    sku = f"RESOLVED-{suffix}"
    rows = [
        submitted_offer_row(
            sku=sku,
            name=f"保留商品-{suffix}",
            source_sheet="Sheet1",
            source_row=2,
        ),
        submitted_offer_row(
            sku=sku,
            name=f"排除商品-{suffix}",
            source_sheet="Sheet2",
            source_row=3,
            included=False,
        ),
        submitted_offer_row(
            sku=f"BAD-{suffix}",
            name=f"错误商品-{suffix}",
            source_sheet="Sheet2",
            source_row=4,
            price="错误价格",
        ),
    ]

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={
            "rows_json": json.dumps(rows, ensure_ascii=False),
            "sheet_configs_json": json.dumps(multi_sheet_configs(), ensure_ascii=False),
        },
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(multi_sheet_offer_workbook_content()),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 201
    data = response.json()
    assert data["status"] == "PARTIAL"
    assert data["total_rows"] == 2
    assert data["success_rows"] == 1
    assert data["error_rows"] == 1
    assert data["errors"][0]["sheet"] == "Sheet2"
    assert data["errors"][0]["row"] == 4
    offers = authenticated_client.get("/api/offers", params={"q": sku}).json()
    assert len(offers) == 1
    assert offers[0]["supplier_sku_code"] == sku


def test_partial_import_persists_only_failed_rows_for_retry(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    successful_sku = f"RETRY-OK-{suffix}"
    failed_sku = f"RETRY-BAD-{suffix}"
    rows = [
        submitted_offer_row(
            sku=successful_sku,
            name=f"成功商品-{suffix}",
            source_sheet="Sheet1",
            source_row=2,
        ),
        submitted_offer_row(
            sku=failed_sku,
            name=f"失败商品-{suffix}",
            source_sheet="Sheet1",
            source_row=3,
            price="错误价格",
        ),
    ]

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps(rows, ensure_ascii=False)},
        files={"file": ("retry.csv", BytesIO(b"ignored"), "text/csv")},
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["status"] == "PARTIAL"
    assert payload["retryable"] is True
    assert "retry_rows" not in payload
    with SessionLocal() as db:
        job = db.get(ImportJob, payload["id"])
        assert job is not None
        assert len(job.retry_rows) == 1
        assert job.retry_rows[0]["values"]["supplier_sku"] == failed_sku


def test_retry_import_creates_linked_job_and_processes_saved_failed_rows(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    retry_sku = f"RETRY-SAVED-{suffix}"
    with SessionLocal() as db:
        supplier_user = db.scalar(select(User).where(User.email == SUPPLIER_EMAIL))
        assert supplier_user is not None
        original = ImportJob(
            organization_id=supplier_user.organization_id,
            file_name="retry-source.xlsx",
            import_type="SMART_PRODUCT_OFFER",
            status=ImportStatus.FAILED.value,
            total_rows=1,
            error_rows=1,
            retry_rows=[
                submitted_offer_row(
                    sku=retry_sku,
                    name=f"重试商品-{suffix}",
                    source_sheet="Sheet1",
                    source_row=9,
                )
            ],
            retry_row_count=1,
        )
        db.add(original)
        db.commit()
        original_id = original.id

    response = authenticated_client.post(f"/api/imports/{original_id}/retry")

    assert response.status_code == 201
    payload = response.json()
    assert payload["status"] == "COMPLETED"
    assert payload["success_rows"] == 1
    assert payload["retry_of_id"] == original_id
    offers = authenticated_client.get("/api/offers", params={"q": retry_sku}).json()
    assert [offer["supplier_sku_code"] for offer in offers] == [retry_sku]

    repeated = authenticated_client.post(f"/api/imports/{original_id}/retry")
    assert repeated.status_code == 409
    assert repeated.json()["detail"] == "该导入任务的失败数据已经重试"
    history = authenticated_client.get("/api/imports").json()
    original_view = next(job for job in history if job["id"] == original_id)
    assert original_view["retryable"] is False


def test_retry_import_rejects_foreign_or_snapshotless_jobs(
    authenticated_client: TestClient,
) -> None:
    with SessionLocal() as db:
        supplier_user = db.scalar(select(User).where(User.email == SUPPLIER_EMAIL))
        assert supplier_user is not None
        other_supplier = Organization(
            code=f"OTHER-{uuid4().hex[:8]}",
            name="其他供应商",
            organization_type=OrganizationType.SUPPLIER.value,
        )
        db.add(other_supplier)
        db.flush()
        foreign_job = ImportJob(
            organization_id=other_supplier.id,
            file_name="foreign.xlsx",
            status=ImportStatus.FAILED.value,
            retry_rows=[
                submitted_offer_row(
                    sku="FOREIGN",
                    name="其他",
                    source_sheet="S",
                    source_row=2,
                )
            ],
            retry_row_count=1,
        )
        old_job = ImportJob(
            organization_id=supplier_user.organization_id,
            file_name="old.xlsx",
            import_type="SMART_PRODUCT_OFFER",
            status=ImportStatus.FAILED.value,
            error_rows=1,
        )
        db.add_all([foreign_job, old_job])
        db.commit()
        foreign_id = foreign_job.id
        old_id = old_job.id

    assert authenticated_client.post(f"/api/imports/{foreign_id}/retry").status_code == 404
    response = authenticated_client.post(f"/api/imports/{old_id}/retry")
    assert response.status_code == 409
    assert response.json()["detail"] == "该导入任务未保留失败数据，请重新选择原文件"


@pytest.mark.parametrize(
    ("overrides", "detail"),
    [
        ({"source_sheet": 123}, "目标列表第 1 行来源工作表格式不正确"),
        ({"source_row": "2"}, "目标列表第 1 行来源行号格式不正确"),
        ({"included": "true"}, "目标列表第 1 行是否导入格式不正确"),
    ],
)
def test_final_import_rejects_invalid_source_metadata(
    authenticated_client: TestClient,
    overrides: dict[str, object],
    detail: str,
) -> None:
    row = submitted_offer_row(
        sku=f"META-{uuid4().hex[:8]}",
        name="元数据测试商品",
        source_sheet="Sheet1",
        source_row=2,
    )
    row.update(overrides)

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={
            "rows_json": json.dumps([row], ensure_ascii=False),
            "sheet_configs_json": json.dumps(multi_sheet_configs(), ensure_ascii=False),
        },
        files={
            "file": (
                "supplier.xlsx",
                BytesIO(multi_sheet_offer_workbook_content()),
                XLSX_MIME,
            )
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == detail


def test_final_excel_import_requires_sheet_configs_for_nested_rows(
    authenticated_client: TestClient,
) -> None:
    row = submitted_offer_row(
        sku=f"NO-CONFIG-{uuid4().hex[:8]}",
        name="缺配置商品",
        source_sheet="Sheet1",
        source_row=2,
    )

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps([row], ensure_ascii=False)},
        files={"file": ("supplier.xlsx", BytesIO(multi_sheet_offer_workbook_content()), XLSX_MIME)},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "Excel 多工作表导入缺少工作表配置"


def test_final_excel_import_rejects_source_sheet_not_selected(
    authenticated_client: TestClient,
) -> None:
    row = submitted_offer_row(
        sku=f"UNSELECTED-{uuid4().hex[:8]}",
        name="未选工作表商品",
        source_sheet="Sheet2",
        source_row=3,
    )

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={
            "rows_json": json.dumps([row], ensure_ascii=False),
            "sheet_configs_json": json.dumps([multi_sheet_configs()[0]], ensure_ascii=False),
        },
        files={"file": ("supplier.xlsx", BytesIO(multi_sheet_offer_workbook_content()), XLSX_MIME)},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "目标列表第 1 行来源工作表未被选择：Sheet2"


def test_final_excel_import_accepts_verified_selected_sheet_source(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    sku = f"VERIFIED-{suffix}"
    row = submitted_offer_row(
        sku=sku,
        name=f"已验证商品-{suffix}",
        source_sheet="Sheet1",
        source_row=2,
    )

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={
            "rows_json": json.dumps([row], ensure_ascii=False),
            "sheet_configs_json": json.dumps([multi_sheet_configs()[0]], ensure_ascii=False),
        },
        files={"file": ("supplier.xlsx", BytesIO(multi_sheet_offer_workbook_content()), XLSX_MIME)},
    )

    assert response.status_code == 201
    assert response.json()["success_rows"] == 1
    offers = authenticated_client.get("/api/offers", params={"q": sku}).json()
    assert offers[0]["supplier_sku_code"] == sku


def test_dashboard_requires_login(client: TestClient) -> None:
    client.post("/api/auth/logout")
    response = client.get("/api/dashboard")
    assert response.status_code == 401


def test_login_and_dashboard(client: TestClient) -> None:
    bad = client.post(
        "/api/auth/login",
        json={"email": SUPPLIER_EMAIL, "password": "wrong-password"},
    )
    assert bad.status_code == 401

    ok = client.post(
        "/api/auth/login",
        json={"email": SUPPLIER_EMAIL, "password": SUPPLIER_PASSWORD},
    )
    assert ok.status_code == 200
    assert ok.json()["user"].get("organization_type") == "SUPPLIER"
    dashboard = client.get("/api/dashboard")
    assert dashboard.status_code == 200
    assert len(dashboard.json()["metrics"]) == 4


def test_supplier_application(client: TestClient) -> None:
    unique = uuid4().hex[:8]
    response = client.post(
        "/api/public/applications",
        json={
            "company_name": f"测试供应商-{unique}",
            "company_type": "生产工厂",
            "province": "浙江省",
            "city": "宁波市",
            "contact_name": "测试联系人",
            "phone": "13800138000",
            "email": f"supplier-{unique}@example.com",
            "categories": ["工业自动化"],
            "cooperation_modes": ["B2B外贸"],
            "supports_dropshipping": True,
        },
    )
    assert response.status_code == 201
    data = response.json()
    assert data["application_no"].startswith("SUP-")
    assert data["status"] == "PENDING"


@pytest.mark.parametrize(
    "changes",
    [
        {"price": "49.9000"},
        {"stock_qty": 321},
        {"moq": 25},
        {"lead_time_days": 14},
    ],
)
def test_create_product_and_offer(
    authenticated_client: TestClient,
    changes: dict[str, object],
) -> None:
    suffix = uuid4().hex[:8]
    product_response = authenticated_client.post(
        "/api/products",
        json={
            "name": f"测试接近传感器-{suffix}",
            "brand": "TEST",
            "model": suffix,
            "category": "工业自动化/接近传感器",
            "attributes": {"输出": "PNP NO"},
            "status": "ACTIVE",
        },
    )
    assert product_response.status_code == 201
    product_id = product_response.json()["id"]

    offer_response = authenticated_client.post(
        "/api/offers",
        json={
            "product_id": product_id,
            "supplier_sku_code": f"SKU-{suffix}",
            "price": 45.8,
            "currency": "CNY",
            "moq": 10,
            "stock_qty": 120,
            "lead_time_days": 3,
            "fulfillment_mode": "PURCHASE",
            "status": "ACTIVE",
        },
    )
    assert offer_response.status_code == 201
    created = offer_response.json()
    supplier_sku_id = created["supplier_sku_id"]
    assert created["supplier_sku_code"] == f"SKU-{suffix}"
    assert created["brand"] == "TEST"
    assert created["brand_id"] == product_response.json()["brand_id"]
    assert created["stock_qty"] == 120

    updated = authenticated_client.patch(
        f"/api/offers/{created['id']}",
        json=changes,
    ).json()

    assert updated["supplier_sku_id"] == supplier_sku_id
    assert updated["supplier_sku_code"] == f"SKU-{suffix}"


def test_product_offset_pages_are_stable_and_filter_before_pagination(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    active_ids = []
    for index in range(3):
        response = authenticated_client.post(
            "/api/products",
            json={
                "name": f"分页商品-{suffix}-{index}",
                "brand": "TEST",
                "model": f"PAGE-{suffix}-{index}",
                "category": "分页测试",
                "status": "ACTIVE",
            },
        )
        assert response.status_code == 201
        active_ids.append(response.json()["id"])
    excluded = authenticated_client.post(
        "/api/products",
        json={
            "name": f"分页商品-{suffix}-excluded",
            "brand": "TEST",
            "model": f"PAGE-{suffix}-excluded",
            "category": "分页测试",
            "status": "DRAFT",
        },
    ).json()

    with SessionLocal() as db:
        for product_id in active_ids:
            db.get(Product, product_id).updated_at = datetime(2000, 1, 1, tzinfo=timezone.utc)
        db.get(Product, excluded["id"]).updated_at = datetime(2001, 1, 1, tzinfo=timezone.utc)
        db.commit()

    first = authenticated_client.get(
        "/api/products",
        params={"q": f"分页商品-{suffix}", "status": "ACTIVE", "limit": 2, "offset": 0},
    )
    second = authenticated_client.get(
        "/api/products",
        params={"q": f"分页商品-{suffix}", "status": "ACTIVE", "limit": 2, "offset": 2},
    )

    assert first.status_code == 200
    assert second.status_code == 200
    assert isinstance(first.json(), list)
    page_ids = [item["id"] for item in first.json() + second.json()]
    assert page_ids == sorted(active_ids, reverse=True)
    assert len(page_ids) == len(set(page_ids)) == 3


def test_offer_offset_pages_are_stable_and_filter_before_pagination(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    product = authenticated_client.post(
        "/api/products",
        json={
            "name": f"报价分页商品-{suffix}",
            "brand": "TEST",
            "model": f"OFFER-PAGE-{suffix}",
            "category": "分页测试",
            "status": "ACTIVE",
        },
    ).json()
    active_ids = []
    for index in range(3):
        response = authenticated_client.post(
            "/api/offers",
            json={
                "product_id": product["id"],
                "supplier_sku_code": f"OFFSET-{suffix}-{index}",
                "price": 10 + index,
                "status": "ACTIVE",
            },
        )
        assert response.status_code == 201
        active_ids.append(response.json()["id"])
    excluded = authenticated_client.post(
        "/api/offers",
        json={
            "product_id": product["id"],
            "supplier_sku_code": f"OFFSET-{suffix}-excluded",
            "price": 99,
            "status": "PAUSED",
        },
    ).json()

    with SessionLocal() as db:
        for offer_id in active_ids:
            db.get(SupplierOffer, offer_id).updated_at = datetime(2000, 1, 1, tzinfo=timezone.utc)
        db.get(SupplierOffer, excluded["id"]).updated_at = datetime(2001, 1, 1, tzinfo=timezone.utc)
        db.commit()

    first = authenticated_client.get(
        "/api/offers",
        params={"q": f"OFFSET-{suffix}", "status": "ACTIVE", "limit": 2, "offset": 0},
    )
    second = authenticated_client.get(
        "/api/offers",
        params={"q": f"OFFSET-{suffix}", "status": "ACTIVE", "limit": 2, "offset": 2},
    )

    assert first.status_code == 200
    assert second.status_code == 200
    assert isinstance(first.json(), list)
    page_ids = [item["id"] for item in first.json() + second.json()]
    assert page_ids == sorted(active_ids, reverse=True)
    assert len(page_ids) == len(set(page_ids)) == 3


def test_offer_contract_rejects_legacy_supplier_sku_payload(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    product = authenticated_client.post(
        "/api/products",
        json={
            "name": f"旧工作台兼容商品-{suffix}",
            "brand": "TEST",
            "model": suffix,
            "category": "工业自动化",
            "status": "ACTIVE",
        },
    ).json()

    response = authenticated_client.post(
        "/api/offers",
        json={
            "product_id": product["id"],
            "supplier_sku_code": f"FINAL-WORKBENCH-{suffix}",
            "supplier_sku": f"LEGACY-WORKBENCH-{suffix}",
            "price": "18.5000",
        },
    )

    assert response.status_code == 422
    assert any(error["loc"][-1] == "supplier_sku" for error in response.json()["detail"])


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("supplier_sku", "LEGACY-PATCH"),
        ("supplier_sku_id", str(uuid4())),
        ("supplier_sku_code", "REPLACEMENT-PATCH"),
        ("unknown", "unexpected"),
    ],
)
def test_offer_patch_rejects_identity_and_unknown_fields_without_changes(
    authenticated_client: TestClient,
    field: str,
    value: str,
) -> None:
    suffix = uuid4().hex[:8]
    product = authenticated_client.post(
        "/api/products",
        json={
            "name": f"报价身份不可变商品-{suffix}",
            "brand": "TEST",
            "model": suffix,
            "category": "工业自动化",
            "status": "ACTIVE",
        },
    ).json()
    created = authenticated_client.post(
        "/api/offers",
        json={
            "product_id": product["id"],
            "supplier_sku_code": f"IMMUTABLE-{suffix}",
            "price": "18.5000",
        },
    )
    assert created.status_code == 201
    original = created.json()

    response = authenticated_client.patch(
        f"/api/offers/{original['id']}",
        json={"price": "19.5000", field: value},
    )

    assert response.status_code == 422
    assert any(error["loc"][-1] == field for error in response.json()["detail"])
    unchanged = authenticated_client.get(
        "/api/offers", params={"q": original["supplier_sku_code"]}
    ).json()
    assert len(unchanged) == 1
    assert unchanged[0]["supplier_sku_id"] == original["supplier_sku_id"]
    assert unchanged[0]["supplier_sku_code"] == original["supplier_sku_code"]
    assert unchanged[0]["price"] == original["price"]


def test_offer_response_uses_only_stable_supplier_sku_fields(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    product = authenticated_client.post(
        "/api/products",
        json={
            "name": f"最终报价契约商品-{suffix}",
            "brand": "TEST",
            "model": suffix,
            "category": "工业自动化",
            "status": "ACTIVE",
        },
    ).json()
    response = authenticated_client.post(
        "/api/offers",
        json={
            "product_id": product["id"],
            "supplier_sku_code": f"FINAL-WORKBENCH-{suffix}",
            "price": "18.5000",
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["supplier_sku_code"] == f"FINAL-WORKBENCH-{suffix}"
    assert str(UUID(body["supplier_sku_id"])) == body["supplier_sku_id"]
    assert "supplier_sku" not in body
    assert body["brand"] == "TEST"


def test_offer_list_uses_only_current_tenant_active_cooperation_mode(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    shared_prefix = f"TENANT-MODE-{suffix}"
    with SessionLocal() as db:
        supplier = db.scalar(
            select(Organization).where(Organization.code == "TEST-SUPPLIER")
        )
        assert supplier is not None
        other_supplier = Organization(
            code=f"OTHER-MODE-{suffix}",
            name=f"其他租户-{suffix}",
            organization_type=OrganizationType.SUPPLIER.value,
        )
        brand = Brand(
            code=f"MODE-{suffix}",
            name=f"合作模式隔离品牌-{suffix}",
            normalized_name=f"合作模式隔离品牌-{suffix}".lower(),
            aliases=[],
            status=CatalogStatus.ACTIVE.value,
        )
        db.add_all([other_supplier, brand])
        db.flush()
        own_product = Product(
            created_by_organization_id=supplier.id,
            brand_id=brand.id,
            name=f"本租户商品-{suffix}",
            category="工业自动化",
            status="ACTIVE",
        )
        other_product = Product(
            created_by_organization_id=other_supplier.id,
            brand_id=brand.id,
            name=f"其他租户商品-{suffix}",
            category="工业自动化",
            status="ACTIVE",
        )
        db.add_all([own_product, other_product])
        db.flush()
        cooperations = [
            SupplierBrandCooperation(
                supplier_id=supplier.id,
                brand_id=brand.id,
                commercial_mode=CommercialMode.SELF_PURCHASE.value,
                status=CatalogStatus.INACTIVE.value,
            ),
            SupplierBrandCooperation(
                supplier_id=supplier.id,
                brand_id=brand.id,
                commercial_mode=CommercialMode.B2B.value,
                status=CatalogStatus.ACTIVE.value,
            ),
            SupplierBrandCooperation(
                supplier_id=other_supplier.id,
                brand_id=brand.id,
                commercial_mode=CommercialMode.JOINT_OPERATION.value,
                status=CatalogStatus.ACTIVE.value,
            ),
        ]
        own_sku = SupplierSku(
            supplier_id=supplier.id,
            brand_id=brand.id,
            product_id=own_product.id,
            supplier_sku_code=f"{shared_prefix}-OWN",
            status=CatalogStatus.ACTIVE.value,
        )
        other_sku = SupplierSku(
            supplier_id=other_supplier.id,
            brand_id=brand.id,
            product_id=other_product.id,
            supplier_sku_code=f"{shared_prefix}-OTHER",
            status=CatalogStatus.ACTIVE.value,
        )
        db.add_all([*cooperations, own_sku, other_sku])
        db.flush()
        own_offer = SupplierOffer(
            organization_id=supplier.id,
            product_id=own_product.id,
            supplier_sku_id=own_sku.id,
            price="10.0000",
            currency="CNY",
            moq=1,
            stock_qty=0,
            lead_time_days=3,
            fulfillment_mode="PURCHASE",
            status="ACTIVE",
        )
        other_offer = SupplierOffer(
            organization_id=other_supplier.id,
            product_id=other_product.id,
            supplier_sku_id=other_sku.id,
            price="20.0000",
            currency="CNY",
            moq=1,
            stock_qty=0,
            lead_time_days=3,
            fulfillment_mode="PURCHASE",
            status="ACTIVE",
        )
        db.add_all([own_offer, other_offer])
        db.commit()
        own_offer_id = own_offer.id
        other_offer_id = other_offer.id

    response = authenticated_client.get("/api/offers", params={"q": shared_prefix})

    assert response.status_code == 200
    assert [(item["id"], item["commercial_mode"]) for item in response.json()] == [
        (own_offer_id, CommercialMode.B2B.value)
    ]
    assert all(item["id"] != other_offer_id for item in response.json())


def test_supplier_workbench_uses_normalized_catalog_controls(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.get("/app")

    assert response.status_code == 200
    assert '<select id="product-brand" class="select" name="brand" required>' in response.text
    assert 'id="offer-supplier-sku-id"' in response.text
    assert 'id="offer-price-label"' in response.text
    assert 'name="commercial_mode"' not in response.text
    assert 'name="supplier_sku_id"' not in response.text


def test_product_requires_an_assigned_non_empty_brand(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    missing = authenticated_client.post(
        "/api/products",
        json={
            "name": f"缺少品牌商品-{suffix}",
            "category": "工业自动化",
        },
    )
    assert missing.status_code == 422

    unassigned = authenticated_client.post(
        "/api/products",
        json={
            "name": f"未分配品牌商品-{suffix}",
            "brand": f"UNASSIGNED-{suffix}",
            "category": "工业自动化",
        },
    )
    assert unassigned.status_code == 400
    assert unassigned.json()["detail"] == "brand is not assigned to supplier"


def test_supplier_catalog_lists_only_active_assigned_brands(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.get(
        "/api/supplier-catalog/brand-cooperations"
    )

    assert response.status_code == 200
    by_code = {item["brand_code"]: item for item in response.json()}
    test_brand = by_code["TEST-BRAND"]
    assert str(UUID(test_brand["brand_id"])) == test_brand["brand_id"]
    assert {key: value for key, value in test_brand.items() if key != "brand_id"} == {
        "brand_code": "TEST-BRAND",
        "brand_name": "TEST",
        "commercial_mode": "SELF_PURCHASE",
        "status": "ACTIVE",
    }
    assert by_code["TEST-BRAND-ALT"]["commercial_mode"] == "B2B"


def test_supplier_updates_only_its_own_brand_commercial_mode(
    authenticated_client: TestClient,
) -> None:
    with SessionLocal() as db:
        supplier = db.scalar(
            select(Organization).where(Organization.code == "TEST-SUPPLIER")
        )
        assert supplier is not None
        own_brand = Brand(
            code=f"MODE-OWN-{uuid4().hex[:8]}",
            name=f"自有模式品牌-{uuid4().hex[:8]}",
            normalized_name=f"own-{uuid4().hex}",
            aliases=[],
            status=CatalogStatus.ACTIVE.value,
        )
        other_supplier = Organization(
            code=f"MODE-OTHER-{uuid4().hex[:8]}",
            name="模式隔离供应商",
            organization_type=OrganizationType.SUPPLIER.value,
        )
        other_brand = Brand(
            code=f"MODE-BRAND-{uuid4().hex[:8]}",
            name=f"其他供应商品牌-{uuid4().hex[:8]}",
            normalized_name=f"other-{uuid4().hex}",
            aliases=[],
            status=CatalogStatus.ACTIVE.value,
        )
        db.add_all([own_brand, other_supplier, other_brand])
        db.flush()
        db.add_all([
            SupplierBrandCooperation(
                supplier_id=supplier.id,
                brand_id=own_brand.id,
                commercial_mode=None,
                status=CatalogStatus.ACTIVE.value,
            ),
            SupplierBrandCooperation(
                supplier_id=other_supplier.id,
                brand_id=other_brand.id,
                commercial_mode=None,
                status=CatalogStatus.ACTIVE.value,
            ),
        ])
        db.commit()
        own_brand_id = own_brand.id
        other_brand_id = other_brand.id

    updated = authenticated_client.put(
        f"/api/supplier-catalog/brand-cooperations/{own_brand_id}",
        json={"commercial_mode": "JOINT_OPERATION"},
    )
    assert updated.status_code == 200
    assert updated.json()["brand_id"] == own_brand_id
    assert updated.json()["commercial_mode"] == "JOINT_OPERATION"

    forbidden = authenticated_client.put(
        f"/api/supplier-catalog/brand-cooperations/{other_brand_id}",
        json={"commercial_mode": "B2B"},
    )
    assert forbidden.status_code == 404


def test_supplier_workbench_consolidates_products_into_product_offers(
    authenticated_client: TestClient,
) -> None:
    page = authenticated_client.get("/app")

    assert page.status_code == 200
    assert 'data-route="products"' not in page.text
    assert 'data-view="products"' not in page.text
    assert "商品主数据" not in page.text
    assert "供应报价" not in page.text
    assert "商品报价" in page.text
    assert 'id="brand-commercial-mode-list"' in page.text


def test_reimporting_supplier_sku_code_reuses_stable_id(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    sku = f"REIMPORT-{suffix}"
    row = submitted_offer_row(
        sku=sku,
        name=f"重导入商品-{suffix}",
        source_sheet="Sheet1",
        source_row=2,
    )

    first = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps([row], ensure_ascii=False)},
        files={"file": ("first.csv", BytesIO(b"source file"), "text/csv")},
    )
    assert first.status_code == 201
    first_offer = authenticated_client.get("/api/offers", params={"q": sku}).json()[0]

    row["values"]["price"] = "12.50"
    second = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps([row], ensure_ascii=False)},
        files={"file": ("second.csv", BytesIO(b"source file"), "text/csv")},
    )
    assert second.status_code == 201
    second_offer = authenticated_client.get("/api/offers", params={"q": sku}).json()[0]

    assert second_offer["supplier_sku_id"] == first_offer["supplier_sku_id"]
    assert second_offer["supplier_sku_code"] == sku
    assert second_offer["price"] == "12.5000"


def test_import_auto_links_unassigned_brand_with_default_a_mode(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    row = submitted_offer_row(
        sku=f"UNASSIGNED-IMPORT-{suffix}",
        name=f"未分配导入商品-{suffix}",
        source_sheet="Sheet1",
        source_row=2,
    )
    row["values"]["brand"] = f"UNASSIGNED-IMPORT-BRAND-{suffix}"

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps([row], ensure_ascii=False)},
        files={"file": ("unassigned.csv", BytesIO(b"source file"), "text/csv")},
    )

    assert response.status_code == 201
    assert response.json()["status"] == "COMPLETED"
    assert response.json()["success_rows"] == 1
    assert response.json()["error_rows"] == 0
    products = authenticated_client.get(
        "/api/products", params={"q": f"未分配导入商品-{suffix}"}
    ).json()
    assert len(products) == 1
    with SessionLocal() as db:
        brand = db.scalar(
            select(Brand).where(
                Brand.normalized_name == row["values"]["brand"].lower()
            )
        )
        assert brand is not None
        cooperation = db.scalar(
            select(SupplierBrandCooperation).where(
                SupplierBrandCooperation.brand_id == brand.id,
                SupplierBrandCooperation.status == CatalogStatus.ACTIVE.value,
            )
        )
        assert cooperation is not None
        assert cooperation.commercial_mode == CommercialMode.SELF_PURCHASE.value


def test_import_history_persists_complete_grouped_failure_counts(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    rows = []
    for index in range(101):
        row = submitted_offer_row(
            sku=f"SUMMARY-{suffix}-{index}",
            name=f"失败汇总商品-{suffix}-{index}",
            source_sheet="Sheet1",
            source_row=index + 2,
        )
        if index < 100:
            row["values"]["category"] = ""
        else:
            row["values"]["price"] = "0"
        rows.append(row)

    created = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps(rows, ensure_ascii=False)},
        files={"file": ("mixed-errors.csv", BytesIO(b"source file"), "text/csv")},
    )

    assert created.status_code == 201
    summaries = {item["reason"]: item for item in created.json()["error_summary"]}
    assert summaries["类目不能为空"]["affected_rows"] == 100
    assert summaries["价格必须大于0"]["affected_rows"] == 1
    history = authenticated_client.get("/api/imports")
    assert history.status_code == 200
    persisted = next(
        item for item in history.json() if item["id"] == created.json()["id"]
    )
    assert persisted["error_summary"] == created.json()["error_summary"]


def test_import_does_not_update_same_code_offer_from_other_tenant(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    sku_code = f"MISMATCHED-TENANT-{suffix}"
    product_name = f"跨租户历史报价商品-{suffix}"
    product_response = authenticated_client.post(
        "/api/products",
        json={
            "name": product_name,
            "brand": "TEST",
            "model": sku_code,
            "category": "工业自动化",
            "status": "ACTIVE",
        },
    )
    assert product_response.status_code == 201

    with SessionLocal() as db:
        supplier = db.scalar(
            select(Organization).where(Organization.code == "TEST-SUPPLIER")
        )
        brand = db.scalar(select(Brand).where(Brand.code == "TEST-BRAND"))
        assert supplier is not None
        assert brand is not None

        other_supplier = Organization(
            code=f"MISMATCHED-OFFER-{suffix}",
            name=f"跨租户报价组织-{suffix}",
            organization_type=OrganizationType.SUPPLIER.value,
        )
        db.add(other_supplier)
        db.flush()
        other_product = Product(
            created_by_organization_id=other_supplier.id,
            brand_id=brand.id,
            name=product_name,
            model=sku_code,
            category="工业自动化",
            status="ACTIVE",
        )
        db.add(other_product)
        db.flush()
        other_supplier_sku = SupplierSku(
            supplier_id=other_supplier.id,
            brand_id=brand.id,
            product_id=other_product.id,
            variant_id=None,
            supplier_sku_code=sku_code,
            status="ACTIVE",
        )
        db.add(other_supplier_sku)
        db.flush()
        other_offer = SupplierOffer(
            organization_id=other_supplier.id,
            product_id=other_product.id,
            supplier_sku_id=other_supplier_sku.id,
            price="99.0000",
            currency="CNY",
            moq=1,
            stock_qty=0,
            lead_time_days=3,
            fulfillment_mode="PURCHASE",
            status="ACTIVE",
        )
        db.add(other_offer)
        db.commit()
        other_offer_id = other_offer.id

    row = submitted_offer_row(
        sku=sku_code,
        name=product_name,
        source_sheet="Sheet1",
        source_row=2,
        price="12.5000",
    )
    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps([row], ensure_ascii=False)},
        files={"file": ("mismatched.csv", BytesIO(b"source file"), "text/csv")},
    )

    assert response.status_code == 201
    assert response.json()["status"] == "COMPLETED"
    assert response.json()["success_rows"] == 1
    assert response.json()["error_rows"] == 0
    with SessionLocal() as db:
        other_offer = db.get(SupplierOffer, other_offer_id)
        assert other_offer is not None
        assert str(other_offer.price) == "99.0000"


def test_offers_can_filter_by_brand(authenticated_client: TestClient) -> None:
    suffix = uuid4().hex[:8]
    brands = ["TEST", "TEST ALT"]
    for index, brand in enumerate(brands, start=1):
        product = authenticated_client.post(
            "/api/products",
            json={
                "name": f"品牌筛选商品-{index}-{suffix}",
                "brand": brand,
                "model": f"MODEL-{index}-{suffix}",
                "category": "工业自动化",
                "status": "ACTIVE",
            },
        ).json()
        response = authenticated_client.post(
            "/api/offers",
            json={
                "product_id": product["id"],
                "supplier_sku_code": f"BRAND-{index}-{suffix}",
                "price": 20 + index,
                "currency": "CNY",
                "moq": 1,
                "stock_qty": 10,
                "lead_time_days": 3,
                "fulfillment_mode": "PURCHASE",
                "status": "ACTIVE",
            },
        )
        assert response.status_code == 201

    brand_response = authenticated_client.get("/api/offers/brands")
    assert brand_response.status_code == 200
    assert all(brand in brand_response.json() for brand in brands)

    filtered = authenticated_client.get(
        "/api/offers",
        params={"brand": brands[0], "q": f"BRAND-1-{suffix}"},
    )
    assert filtered.status_code == 200
    assert len(filtered.json()) == 1
    assert filtered.json()[0]["brand"] == brands[0]


def test_can_delete_all_brand_data_for_selected_brand(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    created = []
    for index, brand in enumerate((f"DELETE-{suffix}", f"KEEP-{suffix}"), start=1):
        product = authenticated_client.post(
            "/api/products",
            json={
                "name": f"品牌批量删除商品-{index}-{suffix}",
                "brand": brand,
                "model": f"MODEL-{index}-{suffix}",
                "category": "工业自动化",
                "status": "ACTIVE",
            },
        ).json()
        offer = authenticated_client.post(
            "/api/offers",
            json={
                "product_id": product["id"],
                "supplier_sku_code": f"DELETE-BRAND-{index}-{suffix}",
                "price": 20 + index,
                "currency": "CNY",
                "moq": 1,
                "stock_qty": 5,
                "lead_time_days": 3,
                "fulfillment_mode": "PURCHASE",
                "status": "ACTIVE",
            },
        ).json()
        created.append((product, offer))

    deleted = authenticated_client.delete(
        "/api/offers/brand", params={"brand": f"DELETE-{suffix}"}
    )
    assert deleted.status_code == 200
    assert deleted.json()["deleted_count"] == 1
    assert deleted.json()["deleted_products"] == 1
    assert deleted.json()["deleted_supplier_skus"] == 1
    assert deleted.json()["deleted_cooperations"] == 1

    assert authenticated_client.get(
        "/api/offers", params={"brand": f"DELETE-{suffix}"}
    ).json() == []
    kept = authenticated_client.get(
        "/api/offers", params={"brand": f"KEEP-{suffix}"}
    ).json()
    assert len(kept) == 1
    with SessionLocal() as db:
        assert db.get(Product, created[0][0]["id"]) is None
        assert db.get(SupplierOffer, created[0][1]["id"]) is None
        assert db.scalar(
            select(InventorySnapshot.id).where(
                InventorySnapshot.offer_id == created[0][1]["id"]
            )
        ) is None


def test_csv_import_supports_partial_success(authenticated_client: TestClient) -> None:
    suffix = uuid4().hex[:8]
    csv_text = (
        "商品名称,品牌,型号,类目,供应商SKU,采购价,币种,起订量,库存,交期天数,履约模式,状态\n"
        f"有效商品,TEST,{suffix},工业自动化/传感器,IMPORT-{suffix},28.5,CNY,5,60,2,PURCHASE,ACTIVE\n"
        f"错误商品,TEST,BAD-{suffix},工业自动化/传感器,IMPORT-BAD-{suffix},错误价格,CNY,5,60,2,PURCHASE,ACTIVE\n"
    )
    response = authenticated_client.post(
        "/api/imports/product-offers",
        files={"file": ("offers.csv", BytesIO(csv_text.encode("utf-8-sig")), "text/csv")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["status"] == "PARTIAL"
    assert data["success_rows"] == 1
    assert data["error_rows"] == 1


def test_smart_import_preview_uses_description_and_detects_columns(
    authenticated_client: TestClient,
) -> None:
    csv_text = (
        "供应商报价单,,,,\n"
        "产品,货号,含税成本,可售数,交货周期\n"
        "工业传感器,A-001,￥28.50,60,7天\n"
    )
    description = (
        "第2行是表头，产品是商品名称，货号作为供应商SKU，含税成本作为采购价，"
        "全部商品类目统一填工业自动化，币种人民币"
    )
    response = authenticated_client.post(
        "/api/imports/product-offers/preview",
        data={"description": description, "defaults_json": json.dumps({"brand": "TEST"})},
        files={"file": ("supplier.csv", BytesIO(csv_text.encode()), "text/csv")},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["header_row"] == 2
    assert data["total_rows"] == 1
    assert data["valid_rows"] == 1
    row = data["preview_rows"][0]["values"]
    assert row["product_name"] == "工业传感器"
    assert row["supplier_sku"] == "A-001"
    assert row["price"] == "￥28.50"
    assert row["category"] == "工业自动化"
    assert row["stock_qty"] == "60"
    assert row["lead_time_days"] == "7天"


def test_smart_xlsx_import_accepts_supplier_original_sheet(
    authenticated_client: TestClient,
) -> None:
    suffix = uuid4().hex[:8]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成本表"
    sheet.append(["供应商内部价格表"])
    sheet.append([])
    sheet.append(["货品名称", "内部料号", "工厂成本", "可供数量"])
    sheet.append([f"露营灯-{suffix}", f"LAMP-{suffix}", 35.8, 120])
    content = BytesIO()
    workbook.save(content)
    content.seek(0)
    description = (
        "成本表第3行是表头，内部料号作为供应商SKU，工厂成本作为采购价，"
        "全部商品类目为户外用品，币种人民币，默认交期5天"
    )

    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={
            "description": description,
            "defaults_json": json.dumps({"brand": "TEST"}),
        },
        files={
            "file": (
                "supplier.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201
    data = response.json()
    assert data["status"] == "COMPLETED"
    assert data["success_rows"] == 1
    offers = authenticated_client.get(
        "/api/offers", params={"search": f"LAMP-{suffix}"}
    ).json()
    assert offers[0]["category"] == "户外用品"
    assert offers[0]["lead_time_days"] == 5


def test_smart_pdf_preview_extracts_cost_table(authenticated_client: TestClient) -> None:
    content = BytesIO()
    document = SimpleDocTemplate(content, pagesize=A4)
    table = Table(
        [
            ["Product Name", "Item Code", "Factory Cost", "Available Qty"],
            ["Camping Light", "PDF-001", "19.90", "30"],
        ],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    document.build([table])
    content.seek(0)
    description = (
        "Item Code作为供应商SKU，Factory Cost作为采购价，"
        "全部商品类目为户外用品，币种美元"
    )

    response = authenticated_client.post(
        "/api/imports/product-offers/preview",
        data={"description": description, "defaults_json": json.dumps({"brand": "TEST"})},
        files={"file": ("supplier.pdf", content, "application/pdf")},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["sheet_name"] == "PDF 第 1 页"
    assert data["valid_rows"] == 1
    row = data["preview_rows"][0]["values"]
    assert row["product_name"] == "Camping Light"
    assert row["supplier_sku"] == "PDF-001"
    assert row["price"] == "19.90"
    assert row["currency"] == "USD"
    assert row["stock_qty"] == "30"


def test_smart_import_uses_edited_target_rows(authenticated_client: TestClient) -> None:
    suffix = uuid4().hex[:8]
    target_rows = [
        {
            "product_name": f"手动新增商品-{suffix}",
            "brand": "TEST",
            "model": suffix,
            "category": "工业自动化",
            "supplier_sku": f"EDIT-{suffix}",
            "price": "88.50",
            "currency": "CNY",
            "moq": "2",
            "stock_qty": "45",
            "lead_time_days": "6",
            "fulfillment_mode": "PURCHASE",
            "status": "ACTIVE",
        }
    ]
    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps(target_rows, ensure_ascii=False)},
        files={"file": ("edited.csv", BytesIO(b"source file"), "text/csv")},
    )

    assert response.status_code == 201
    data = response.json()
    assert data["total_rows"] == 1
    assert data["success_rows"] == 1
    offers = authenticated_client.get(
        "/api/offers", params={"q": f"EDIT-{suffix}"}
    ).json()
    assert len(offers) == 1
    assert offers[0]["product_name"] == f"手动新增商品-{suffix}"


def _invalid_flat_rows(count: int, *, excluded_index: int | None = None) -> list[dict[str, object]]:
    return [
        {
            "supplier_sku": str(index),
            **({"included": False} if index == excluded_index else {}),
        }
        for index in range(count)
    ]


def _nested_boundary_rows(
    count: int,
    *,
    excluded_index: int | None = None,
    suffix: str = "",
) -> list[dict[str, object]]:
    rows = []
    for index in range(count):
        rows.append(
            {
                "source_sheet": "Sheet1",
                "source_row": index + 2,
                "included": index != excluded_index,
                "values": {
                    "product_name": f"边界商品-{suffix}" if index == 0 else "",
                    "brand": "TEST",
                    "model": f"M1-{suffix}",
                    "category": "测试类目",
                    "supplier_sku": f"BOUNDARY-{suffix}-{index}",
                    "price": "1.00" if index == 0 else "",
                    "currency": "CNY",
                    "moq": "1",
                    "stock_qty": "0",
                    "lead_time_days": "3",
                    "fulfillment_mode": "PURCHASE",
                    "status": "ACTIVE",
                },
            }
        )
    return rows


def _boundary_import_request(
    authenticated_client: TestClient,
    count: int,
    *,
    excluded_index: int | None = None,
):
    suffix = uuid4().hex[:8]
    return authenticated_client.post(
        "/api/imports/product-offers",
        data={
            "rows_json": json.dumps(
                _nested_boundary_rows(
                    count,
                    excluded_index=excluded_index,
                    suffix=suffix,
                ),
                ensure_ascii=False,
                separators=(",", ":"),
            ),
            "sheet_configs_json": json.dumps([multi_sheet_configs()[0]]),
        },
        files={
            "file": (
                "boundary.xlsx",
                BytesIO(multi_sheet_offer_workbook_content("UNIQUE-002")),
                XLSX_MIME,
            )
        },
    )


def test_final_import_accepts_exactly_ten_thousand_included_rows(
    authenticated_client: TestClient,
) -> None:
    response = _boundary_import_request(authenticated_client, 10_000)

    assert response.status_code == 201
    assert response.json()["total_rows"] == 10_000
    assert response.json()["success_rows"] == 1
    assert response.json()["status"] == "PARTIAL"


def test_final_import_rejects_ten_thousand_and_one_included_rows(
    authenticated_client: TestClient,
) -> None:
    response = _boundary_import_request(authenticated_client, 10_001)

    assert response.status_code == 400
    assert response.json()["detail"] == "单次最多导入 10000 行"


def test_final_import_excluded_row_does_not_count_toward_ten_thousand_limit(
    authenticated_client: TestClient,
) -> None:
    response = _boundary_import_request(
        authenticated_client,
        10_001,
        excluded_index=10_000,
    )

    assert response.status_code == 201
    assert response.json()["total_rows"] == 10_000
    assert response.json()["success_rows"] == 1


def test_final_import_rejects_oversized_rows_form_part(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": "x" * (MAX_IMPORT_FORM_PART_SIZE + 1)},
        files={"file": ("boundary.csv", BytesIO(b"source"), "text/csv")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "Part exceeded maximum size of 8192KB."


def test_final_import_rejects_total_multipart_request_above_application_limit(
    authenticated_client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        import_routes,
        "MAX_FINAL_IMPORT_REQUEST_SIZE",
        1_024,
        raising=False,
    )
    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": "[]" + (" " * 600)},
        files={"file": ("boundary.csv", BytesIO(b"x" * 600), "text/csv")},
    )

    assert response.status_code == 413
    assert response.json()["detail"] == "导入请求不能超过 20MB"


def test_upload_reader_uses_bounded_chunks_before_rejecting_large_file() -> None:
    class OversizedUpload:
        filename = "oversized.xlsx"

        def __init__(self) -> None:
            self.remaining = import_routes.MAX_FILE_SIZE + 1
            self.read_sizes: list[int] = []

        async def read(self, size: int = -1) -> bytes:
            self.read_sizes.append(size)
            if size < 0:
                raise AssertionError("upload reader requested the entire file")
            count = min(size, self.remaining)
            self.remaining -= count
            return b"x" * count

    upload = OversizedUpload()

    with pytest.raises(import_routes.HTTPException) as error:
        asyncio.run(import_routes._read_upload(upload))

    assert error.value.status_code == 413
    assert error.value.detail == "文件不能超过 10MB"
    assert upload.read_sizes
    assert all(0 < size <= 1024 * 1024 for size in upload.read_sizes)


def test_legacy_flat_rows_default_source_coordinates(
    authenticated_client: TestClient,
) -> None:
    response = authenticated_client.post(
        "/api/imports/product-offers",
        data={"rows_json": json.dumps(_invalid_flat_rows(1), ensure_ascii=False)},
        files={"file": ("legacy.csv", BytesIO(b"source"), "text/csv")},
    )

    assert response.status_code == 201
    assert response.json()["errors"][0]["sheet"] is None
    assert response.json()["errors"][0]["row"] == 1
